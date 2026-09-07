"""Build the article DOCX and readable Markdown from prose and saved result tables.

Requires requirements-docs.txt. Uses native editable Word equations, not equation images.
Run build_article_figures.py first. The experiment itself is not re-estimated here.
"""
from pathlib import Path
import copy
import csv
import hashlib
import json
import math
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / 'paper'
ART = ROOT / 'artifacts'
LABELS = {'equal_weight':'Equal weight','baseline':'Baseline GMV','pipeline1':'Pipeline 1','pipeline2':'Pipeline 2','pipeline3':'Pipeline 3'}

def read_csv(name):
    with (ART / 'tables' / name).open(newline='') as f:
        return list(csv.DictReader(f))

def pct(x, decimals=2): return f'{100*float(x):.{decimals}f}%'
def num(x, decimals=3): return f'{float(x):.{decimals}f}'
def metrics_table(name, columns):
    rows = []
    for r in read_csv(name):
        values = []
        for _,k,mode,d in columns:
            if k=='Average Portfolio Quality':
                split=name.removesuffix('_metrics.csv')
                path=ART/'results'/split/r['']/'diagnostics.csv'
                with path.open(newline='') as f:
                    quality=[float(row['portfolio_quality']) for row in csv.DictReader(f)]
                verified=bool(quality) and all(math.isfinite(q) for q in quality) and math.isclose(sum(quality)/len(quality),float(r[k]),rel_tol=0,abs_tol=1e-9)
                values.append(num(r[k],d) if verified else 'N/V')
            else: values.append(pct(r[k],d) if mode=='pct' else num(r[k],d))
        rows.append([LABELS[r['']]] + values)
    return [['Strategy']+[h for h,_,_,_ in columns]]+rows

def tables():
    common = [('CAGR','CAGR','pct',2),('Volatility','Annualized Volatility','pct',2),('Sharpe','Sharpe','num',3),('CVaR95','CVaR 95%','pct',3),('Max drawdown','Maximum Drawdown','pct',2)]
    out = {
        'splits': [['Split','Price dates','Sessions','Rows','Role'],['Train','2000-01-03 to 2016-12-30','4,277','213,850','Historical estimation'],['Validation','2017-01-03 to 2020-12-31','1,007','50,350','New parameter selection'],['Test','2021-01-04 to 2026-02-20','1,289','64,450','Common evaluation']],
        'quality': [['Domain','Ratio definitions','Favorable direction'],['Profitability','Net income / assets; operating income / revenue','Higher; higher'],['Balance sheet','Equity / assets; cash / assets; liabilities / assets','Higher; higher; lower'],['Cash generation','CFO / assets; (net income − CFO) / assets; (CFO − capital expenditure) / assets','Higher; lower; higher'],['Liquidity','Current assets / current liabilities','Higher']],
        'features': [['Block','Inputs','Examples'],['Price and regimes','21','Momentum, reversal, volatility, drawdown, beta, residual risk, covariance diagonals, HMM probability, market correlation'],['Fundamentals','13','Nine ratios, revenue growth, log assets, quality score and coverage'],['Activity and reporting','12','Volume, illiquidity proxy, range, filing age and lag, amendments, missingness, month seasonality'],['Categorical identity','1 category','Ticker; total feature count is 47']],
        'validation': metrics_table('validation_metrics.csv',common),
        'test': metrics_table('test_metrics.csv',common),
        'implementation': metrics_table('test_metrics.csv',[('Sortino','Sortino','num',3),('CVaR99','CVaR 99%','pct',3),('Turnover / year','Annualized Turnover','num',2),('Effective assets','Average Effective Assets','num',2),('Quality exposure','Average Portfolio Quality','num',3)]),
        'prediction': [['Forecast','Stock months','Months','Log risk RMSE','Mean rank IC']] + [[r['model'],r['rows'],r['months'],num(r['log_risk_RMSE']),num(r['mean_monthly_rank_IC'])] for r in read_csv('test_prediction_metrics.csv')],
        'ablation': [['CatBoost features','Log risk RMSE','Mean rank IC','Median rank IC']] + [[{'price_regime':'Prices and regimes','plus_fundamentals':'Plus fundamentals','full':'Full information'}[r['']],num(r['log_risk_RMSE']),num(r['mean_monthly_rank_IC']),num(r['median_monthly_rank_IC'])] for r in read_csv('test_forecast_feature_ablation.csv')],
        'settings': [['Component','Reported setting'],['Universe and schedule','50 fixed stocks; monthly formation; long only; fully invested; 15% cap'],['Covariance','252 prior sessions; Ledoit–Wolf scaled-identity shrinkage'],['HMM history','2,520 observations maximum; 756 minimum; two Gaussian states'],['HMM fitting','Five starts; 500 iterations; tolerance 10^-4; seed 20260902 plus restart'],['HMM signals','Market return; volatility 20; drawdown 63; correlation 60 sessions'],['State covariance','252 sessions; effective-sample pooling; 21-step average probability'],['Quality','Four domains; nine ratios; lambda 1; missing rank 0.5'],['CatBoost','Depth 4; 300 trees; learning rate 0.04; L2 10; eta 0.25'],['ML chronology','Annual expanding refits; 21-session labels; strict maturity purge'],['ML reproducibility','Seed 20260906; two threads; has_time enabled'],['Costs','10 bps base; two-sided turnover; initial entry; no terminal liquidation'],['Bootstrap','2,000 paired draws; blocks 10, 20 and 60; seed 20260906']]
    }
    pairs=[('baseline','pipeline1'),('pipeline1','pipeline2'),('pipeline2','pipeline3')]
    out['bootstrap']=[['Change','Metric','Difference','95% interval']]
    for a,b in pairs:
        for r in read_csv('test_paired_bootstrap.csv'):
            if r['reference']==a and r['experiment']==b and r['block_length']=='20':
                factor=1 if r['metric']=='Sharpe' else 100
                label={'Annualized Volatility':'Volatility pp','CVaR 95%':'CVaR95 pp','Maximum Drawdown':'Drawdown pp','Sharpe':'Sharpe'}[r['metric']]
                vals=[float(r[k])*factor for k in ['difference','lower_95','upper_95']]
                change=f'{LABELS[b]} − {LABELS[a]}'
                out['bootstrap'].append([change,label,f'{vals[0]:+.4f}',f'[{vals[1]:+.4f}, {vals[2]:+.4f}]'])
    cost=read_csv('test_cost_sensitivity.csv')
    out['costs']=[['Strategy','0 bps','10 bps','25 bps','50 bps']]
    for key,label in LABELS.items():
        out['costs'].append([label]+[pct(next(r['CAGR'] for r in cost if r['strategy']==key and int(r['cost_bps'])==bps)) for bps in [0,10,25,50]])
    return out

# Native Office Math building blocks. Fractions, scripts, delimiters and sums
# remain structured and editable in Word. The accompanying Markdown retains TeX.
def node(tag, **attrs):
    e=OxmlElement('m:'+tag)
    for k,v in attrs.items():e.set(qn('m:'+k),str(v))
    return e
def mr(s):
    e=node('r'); props=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Cambria Math');fonts.set(qn('w:hAnsi'),'Cambria Math');props.append(fonts)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'22');props.append(size);e.append(props)
    t=node('t');t.text=str(s);t.set(qn('xml:space'),'preserve');e.append(t);return e
def seq(*args):
    out=[]
    for a in args:
        if isinstance(a,str):out.append(mr(a))
        elif isinstance(a,(list,tuple)):out.extend(seq(*a))
        else:out.append(copy.deepcopy(a))
    return out
def box(tag,*args):
    e=node(tag);e.extend(seq(*args));return e
def sub(a,b):
    e=node('sSub');e.extend([box('e',a),box('sub',b)]);return e
def sup(a,b):
    e=node('sSup');e.extend([box('e',a),box('sup',b)]);return e
def ss(a,b,c):
    e=node('sSubSup');e.extend([box('e',a),box('sub',b),box('sup',c)]);return e
def frac(a,b):
    e=node('f');e.extend([box('num',a),box('den',b)]);return e
def delim(a,l='(',r=')'):
    e=node('d');p=node('dPr');p.extend([node('begChr',val=l),node('endChr',val=r)]);e.extend([p,box('e',a)]);return e
def summ(lo,hi,a):
    e=node('nary');p=node('naryPr');p.extend([node('chr',val='∑'),node('limLoc',val='subSup')])
    if hi=='':p.append(node('supHide',val=1))
    e.extend([p,box('sub',lo),box('sup',hi),box('e',a)]);return e
def sigma(label):return ss('Σ','t',label)
def equations():
    g=sub('γ','s,k');ne=ss('n','k,t','eff');a=sub('a','k,t')
    w=sub('w','t');q=sub('q','t');dm=ss('D','i,t','−');dp=ss('D','i,t','+')
    gross=ss('r','p,t','gross');net=ss('r','p,t','net')
    ge=[sub('L','t'),'≥',sub('VaR̂','α')]
    return {
        'objective':seq(ss('w','t','*'),' = ',sub('arg min','w ∈ W'),' ',sup('w','T'),sub('Σ','t'),'w', '     W = ',delim(['w: ',sup('1','T'),'w = 1,  0 ≤ ',sub('w','i'),' ≤ 0.15'],'{','}')),
        'shrinkage':seq(sigma('LW'),' = ',delim(['1 − ',sub('δ','t')]),sub('S','t'),' + ',sub('δ','t'),sub('μ','t'),'I', '     ',sub('μ','t'),' = ',frac(['tr',delim(sub('S','t'))],'N')),
        'probability':seq(sub('p̄','t'),' = ',frac('1','21'),summ('h = 1','21',[sub('p','t−1'),ss('A','t','h')])),
        'effective':seq(ne,' = ',frac(sup(delim(summ('s','',g)),'2'),summ('s','',sup(g,'2'))),'     ',a,' = ',frac(ne,[ne,' + N'])),
        'regime':seq(sigma('R'),' = ',summ('k = 1','2',[sub('p̄','k,t'),delim([a,sub('S','k,t'),' + ',delim(['1 − ',a]),sigma('LW')],'[',']')])),
        'quality':seq(sub('q','i,t'),' = ',frac('1','4'),summ('d = 1','4',[frac('1',delim(sub('D','d'),'|','|')),summ(['j ∈ ',sub('D','d')],'',sub('R','i,j,t'))])),
        'qualitymatrix':seq(sigma('Q'),' = ',sigma('R'),' + λ ',sub('m','t'),' diag',delim(['1 − ',q]),'     ',sub('m','t'),' = median',delim(['diag',delim(sigma('R'))])),
        'target':seq(sub('y','i,t'),' = log',delim(frac(dp,dm)),'     ',sub('D̂','i,t'),' = ',dm,' exp',delim(['clip',delim([sub('f','t'),delim(sub('x','i,t')),', −3, 3'])])),
        'mlmatrix':seq(sigma('C'),' = ',sigma('Q'),' + η diag',delim(sub('D̂','t'))),
        'accounting':seq(sub('τ','t'),' = ',summ('i','',delim([ss('w','i,t','*'),' − ',ss('w','i,t','−')],'|','|')),'     ',net,' = ',delim(['1 − c',sub('τ','t')]),delim(['1 + ',gross]),' − 1'),
        'drift':seq(gross,' = ',summ('i','',[sub('w','i,t'),sub('r','i,t')]),'     ',ss('w','i,t+1','−'),' = ',frac([sub('w','i,t'),delim(['1 + ',sub('r','i,t')])],['1 + ',gross])),
        'cvar':seq(sub('L','t'),' = −',net,'     ',sub('CVaR̂','α'),' = ',frac(summ('t','',[sub('L','t'),'1',delim(ge,'{','}')]),summ('t','',['1',delim(ge,'{','}')]))),
    }

def inline(p,text):
    # Keep prose variables as editable inline mathematics, including table tolerances.
    symbols = {
        'r_i,t': sub('r','i,t'), 'w_t': sub('w','t'),
        'q_i,t': sub('q','i,t'), 'R_i,j,t': sub('R','i,j,t'),
        'D_d': sub('D','d'), 'm_t': sub('m','t'),
        'D_i,t^-': ss('D','i,t','−'), 'D_i,t^+': ss('D','i,t','+'),
        'w_t^-': ss('w','t','−'), 'w_t^*': ss('w','t','*'),
        'w_i,t': sub('w','i,t'),
    }
    tokens='|'.join(re.escape(s) for s in sorted(symbols,key=len,reverse=True))
    pattern=re.compile(r'(?<!\w)(?:'+tokens+r'|10\^-\d+)(?!\w)')
    for part in re.split(r'(\*\*.*?\*\*)',text):
        bold=part.startswith('**') and part.endswith('**')
        if bold:part=part[2:-2]
        pos=0
        for match in pattern.finditer(part):
            p.add_run(part[pos:match.start()]).bold=bold
            token=match.group();math=node('oMath')
            math.append(copy.deepcopy(symbols[token]) if token in symbols else sup('10','−'+token.split('^-')[1]))
            p._p.append(math);pos=match.end()
        p.add_run(part[pos:]).bold=bold

def hyperlink(p,label,url):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'000000');pr.append(color);r.append(pr)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'19');pr.append(size)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def add_table(doc, data, key):
    widths={
        'splits':[1.8,5.0,1.65,1.85,6.0], 'quality':[3.0,10.0,3.3], 'features':[3.3,2.2,10.8],
        'validation':[3.25,2.0,2.2,1.7,2.1,3.05], 'test':[3.25,2.0,2.2,1.7,2.1,3.05],
        'implementation':[3.1,1.65,1.85,3.0,3.0,2.2], 'prediction':[5.1,2.55,1.65,3.4,3.6],
        'ablation':[6.1,3.4,3.5,3.3], 'bootstrap':[5.5,3.2,2.4,5.2],
        'costs':[5.0,2.825,2.825,2.825,2.825], 'settings':[4.0,12.3]
    }.get(key,[16.3/len(data[0])]*len(data[0]))
    widths=[v*16.3/sum(widths) for v in widths]
    t=doc.add_table(rows=0, cols=len(data[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,wid in zip(t.columns,widths):c.width=Cm(wid)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
        borders.append(e)
    pr.append(borders)
    for ri,row in enumerate(data):
        cells=t.add_row().cells;trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,(cell,text) in enumerate(zip(cells,row)):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for side,val in [('top','65'),('bottom','65'),('left','90'),('right','90')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');mar.append(e)
            cp.append(mar)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2' if ri==0 else 'FFFFFF');cp.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.02
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or key in ['quality','features','settings'] else WD_ALIGN_PARAGRAPH.CENTER
            inline(p,str(text))
            for r in p.runs:r.font.size=Pt(9.2);r.bold=(ri==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def main():
    raw=(PAPER/'manuscript.template.md').read_text();ts=tables();refs=json.loads((PAPER/'references.json').read_text())
    resolved=raw
    for k,rows in ts.items():
        md='\n'.join(['| '+' | '.join(rows[0])+' |','| '+' | '.join(['---']*len(rows[0]))+' |']+['| '+' | '.join(r)+' |' for r in rows[1:]])
        resolved=resolved.replace('{{table:'+k+'}}',md)
    refmd='\n\n'.join(f"[{r['id']}] {r['text']} [Source]({r['url']})." for r in refs)
    resolved=resolved.replace('{{references}}',refmd)
    assert '{{' not in resolved
    (PAPER/'manuscript.md').write_text(resolved)

    d=Document();s=d.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(2.1);s.bottom_margin=Cm(2.1);s.left_margin=Cm(2.35);s.right_margin=Cm(2.35)
    s.footer_distance=Cm(.9)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
        sty=d.styles[name];sty.font.name='Times New Roman';sty.font.color.rgb=RGBColor(0,0,0)
        sty.font.size=Pt(11)
        fonts=sty.element.get_or_add_rPr().get_or_add_rFonts()
        for a in list(fonts.attrib):
            if 'theme' in a.lower():del fonts.attrib[a]
        for a in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+a),'Times New Roman')
    normal=d.styles['Normal'].paragraph_format;normal.space_after=Pt(5);normal.line_spacing=1.12;normal.widow_control=True
    d.styles['Title'].font.size=Pt(19);d.styles['Title'].font.bold=True
    d.styles['Title'].paragraph_format.space_after=Pt(14)
    for name,size,before,after in [('Heading 1',13,13,6),('Heading 2',11.5,10,4)]:
        st=d.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after);st.paragraph_format.keep_with_next=True
    d.styles['Caption'].font.size=Pt(9.5);d.styles['Caption'].font.italic=False;d.styles['Caption'].font.bold=False
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    d.core_properties.title='Incremental information in minimum variance portfolio optimization'
    d.core_properties.subject='Regimes, business quality and CatBoost downside risk'
    d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.keywords='GMV, HMM, fundamentals, CatBoost, CVaR'
    eqs=equations();n_eq=0;eq_key=None;skip_tex=False;in_refs=False
    lines=raw.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('<!-- equation:'):
            eq_key=line.split(':',1)[1].split(' ',1)[0];continue
        if line=='$$':
            if not skip_tex:
                assert eq_key in eqs,eq_key
                if d.paragraphs:d.paragraphs[-1].paragraph_format.keep_with_next=True
                n_eq+=1;p=d.add_paragraph();p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(7);p.paragraph_format.keep_together=True
                p.paragraph_format.tab_stops.add_tab_stop(Cm(16.1),WD_TAB_ALIGNMENT.RIGHT)
                math=node('oMath');math.extend(eqs[eq_key]);p._p.append(math);p.add_run(f'\t({n_eq})')
            skip_tex=not skip_tex;continue
        if skip_tex:continue
        if line.startswith('{{table:'):
            key=line[len('{{table:'):-2];add_table(d,ts[key],key);continue
        if line=='{{references}}':
            for r in refs:
                p=d.add_paragraph();p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1.0;p.paragraph_format.left_indent=Cm(.7);p.paragraph_format.first_line_indent=Cm(-.7)
                rr=p.add_run(f"[{r['id']}] {r['text']} ");rr.font.size=Pt(9.5)
                label=r['url'].replace('https://doi.org/','doi: ') if 'doi.org/' in r['url'] else 'Online source'
                hyperlink(p,label,r['url'])
            continue
        if line.startswith('!['):
            m=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',line);assert m
            image_path=PAPER/m.group(2);assert image_path.exists(),image_path
            p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            shape=p.add_run().add_picture(str(image_path),width=Cm(16.1));shape._inline.docPr.set('descr',m.group(1));continue
        if line.startswith('# '):
            p=d.add_paragraph(line[2:],style='Title');continue
        if line.startswith('### '):d.add_paragraph(line[4:],style='Heading 2');continue
        if line.startswith('## '):d.add_paragraph(line[3:],style='Heading 1');continue
        following=next((s.strip() for s in lines[i:] if s.strip()),'')
        if re.match(r'^Table (?:\d+|A\d+) ',line) and following.startswith('{{table:'):
            p=d.add_paragraph(line,style='Caption');p.runs[0].bold=True;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(8);continue
        if re.match(r'^Figure \d+\.',line):
            p=d.add_paragraph(line,style='Caption');p.paragraph_format.space_after=Pt(10);continue
        p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;inline(p,line)
    assert not skip_tex and n_eq==len(eqs)
    # Remove theme colors and paragraph border residues in every title/heading style.
    for name in ['Title','Heading 1','Heading 2','Heading 3']:
        st=d.styles[name]
        for e in st.element.xpath('.//w:color'):
            for a in list(e.attrib):
                if 'theme' in a:del e.attrib[a]
        for e in st.element.xpath('.//w:pBdr'):e.getparent().remove(e)
    destination=ROOT/'Portfolio optimization.docx';d.save(destination)
    prov={'source_template_sha256':hashlib.sha256(raw.encode()).hexdigest(),'experiment_manifest_sha256':hashlib.sha256((ART/'run_manifest.json').read_bytes()).hexdigest(),'frozen_spec_sha256':hashlib.sha256((ART/'frozen_spec.json').read_bytes()).hexdigest(),'tables':len(ts),'figures':3,'native_equations':n_eq,'references':len(refs),'word_count_markdown':len(re.findall(r'\b[\w-]+\b',resolved)),'docx_sha256':hashlib.sha256(destination.read_bytes()).hexdigest(),'result_table_sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted((ART/'tables').glob('*.csv'))}}
    prov['presentation_source_sha256']={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [Path(__file__),ROOT/'scripts/build_article_figures.py',PAPER/'references.json',PAPER/'manuscript.md',*sorted((PAPER/'figures').glob('*.png'))]}
    prov['experiment_status']='Historical outputs retained; updated implementation requires a new raw-data validation and freeze'
    (PAPER/'article_manifest.json').write_text(json.dumps(prov,indent=2)+'\n')
    print(json.dumps({k:v for k,v in prov.items() if k in ['tables','figures','native_equations','references','word_count_markdown']},indent=2))

if __name__=='__main__':main()
